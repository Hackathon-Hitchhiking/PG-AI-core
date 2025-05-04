import logging
import os

from pptx import Presentation


def dump_attrs(obj, depth=2, prefix=''):
    """Рекурсивно выводит атрибуты объекта и их значения."""
    if depth < 0:
        return
    for attr in dir(obj):
        try:
            value = getattr(obj, attr)
        except Exception:
            continue
        if callable(value):
            continue
        logging.debug(f'{prefix}{attr}: {repr(value)}')
        if not isinstance(value, (str, int, float, bool, bytes, tuple, list, dict, set)):
            dump_attrs(value, depth - 1, prefix + '  ')


def print_caps_info(shape):
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            rPr = run._r.get_or_add_rPr()
            cap = rPr.get('cap')
            print(f"Run text: '{run.text}', cap: {cap}")


def print_all_shapes_caps(slide):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                rPr = run._r.get_or_add_rPr()
                cap = rPr.get('cap')
                print(f"Shape: '{shape.name}', Run: '{run.text}', cap: {cap}")


def main():
    # Настройка логгирования
    logging.basicConfig(filename='debug.log', filemode='w', level=logging.DEBUG, format='%(message)s')

    pptx_path = os.path.join('test_data', 'final_test_1.pptx')
    prs = Presentation(pptx_path)
    slide = prs.slides[0]

    target_shape = None
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text.strip()
        if text == 'эффекты':
            target_shape = shape
            break

    if target_shape is None:
        logging.debug("Shape with text 'эффекты' not found on the first slide.")
        return

    logging.debug('Attributes of the target shape:')
    dump_attrs(target_shape, depth=2)

    # Print capitalization info for all runs in the shape
    print_caps_info(target_shape)

    # Print cap info for all shapes on the slide
    print_all_shapes_caps(slide)


if __name__ == '__main__':
    main()

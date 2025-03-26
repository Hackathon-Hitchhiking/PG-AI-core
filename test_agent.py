from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.dml.color import RGBColor
from typing import List, Dict, Any
import logging
from pathlib import Path
from pptx_manager.chart import ChartManager, ChartShape

# Настройка логирования
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def parse_presentation_charts(file_path: str) -> List[Dict[str, Any]]:
    """Парсинг всех графиков в презентации"""
    try:
        manager = ChartManager(file_path)
        charts = manager.parse_all_charts()
        return [{
            'slide_id': chart.metadata.slide_id,
            'shape_id': chart.metadata.shape_id,
            'chart_data': chart.model_dump(),
            'position': chart.metadata.size
        } for chart in charts]
    except Exception as e:
        logger.error(f"Parsing failed: {str(e)}")
        return []

def recreate_presentation(parsed_data: List[Dict[str, Any]], output_path: str) -> None:
    """Создание новой презентации из распарсенных данных"""
    try:
        prs = Presentation()
        manager = ChartManager(str(Path(output_path).with_suffix('.tmp')))
        
        for chart_info in parsed_data:
            try:
                chart_shape = ChartShape(**chart_info['chart_data'])
                
                # Создаем новый слайд
                slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
                
                # Подготовка данных для создания графика
                chart_data = {
                    'chart_type': chart_shape.chart_type.value,
                    'position': {
                        'left': chart_info['position']['left'],
                        'top': chart_info['position']['top'],
                        'width': chart_info['position']['width'],
                        'height': chart_info['position']['height']
                    },
                    'data': {
                        'categories': chart_shape.data.categories,
                        'series': [{
                            'name': s.name,
                            'values': s.values,
                            'series_type': s.series_type.value,
                            'color': s.color,
                            'marker': s.marker
                        } for s in chart_shape.data.series]
                    },
                    'style': {
                        'title': chart_shape.metadata.title,
                        'legend': chart_shape.metadata.legend.dict() if chart_shape.metadata.legend else None,
                        'axes': {
                            axis: meta.dict()
                            for axis, meta in chart_shape.metadata.axes.items()
                        }
                    }
                }
                
                # Создание графика
                created = manager.create_chart(
                    slide_id=chart_info['slide_id'],
                    chart_data=chart_data
                )
                
                if not created:
                    logger.warning(f"Failed to recreate chart {chart_info['shape_id']}")
                    
            except Exception as e:
                logger.error(f"Error processing chart {chart_info.get('shape_id', '?')}: {str(e)}")
        
        # Сохранение результата
        prs.save(output_path)
        logger.info(f"Presentation saved to {output_path}")
        
    except Exception as e:
        logger.error(f"Critical error: {str(e)}")
        raise

def print_detailed_chart_info(charts_data: List[Dict[str, Any]]) -> None:
    """Вывод детальной информации о графиках"""
    for idx, chart in enumerate(charts_data, 1):
        data = chart['chart_data']
        print(f"\n{'=' * 50}")
        print(f"Chart {idx} (Slide {chart['slide_id']})")
        print(f"Type: {data['chart_type']}")
        print(f"Title: {data['metadata']['title'] or 'No title'}")
        print(f"Position (inches): L:{chart['position']['left']:.1f}, " 
              f"T:{chart['position']['top']:.1f}, "
              f"W:{chart['position']['width']:.1f}, "
              f"H:{chart['position']['height']:.1f}")
        
        print("\nCategories:", data['data']['categories'])
        print("Series:")
        for s in data['data']['series']:
            print(f" - {s['name']}: {len(s['values'])} values")
            if s.get('color'):
                print(f"   Color: {s['color'].get('fill', {}).get('color', 'default')}")

if __name__ == "__main__":
    # Пример использования
    try:
        # Парсинг исходной презентации
        parsed = parse_presentation_charts("test_sources/test_dit.pptx")
        print_detailed_chart_info(parsed)
        
        # Создание новой презентации
        recreate_presentation(parsed, "test_sources/result.pptx")
        logger.info("Operation completed successfully")
        
    except Exception as e:
        logger.error(f"Fatal error: {str(e)}")
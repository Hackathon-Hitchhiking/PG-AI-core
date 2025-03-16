import os
import json
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
import textwrap

from .utils import rgb_to_hex, hex_to_rgb, resolve_chart_type
from .models import Coordinates, FontStyle, TextShape, ImageShape, ChartSeries, ChartShape, Slide, PresentationInfo


class PPTXHandler:
    def __init__(self, template_json_path):
        with open(template_json_path, "r", encoding="utf-8") as f:
            self.template = json.load(f)["slide_template"]

    @staticmethod
    def parse_font(run):
        font = run.font

        size = font.size.pt if font.size else 18

        color = "#000000"
        if font.color:
            if font.color.type == MSO_COLOR_TYPE.RGB and font.color.rgb:
                color = rgb_to_hex(font.color.rgb)
            elif font.color.type == MSO_COLOR_TYPE.SCHEME:
                color = "#000000"

        name = font.name or "Calibri"
        bold = font.bold if font.bold is not None else False
        italic = font.italic if font.italic is not None else False

        return FontStyle(
            name=name,
            size=size,
            color=color,
            bold=bold,
            italic=italic,
        )

    def parse_chart(self, chart) -> tuple:
        """Парсинг данных графика."""
        plot = chart.plots[0]
        categories = [c.label for c in plot.categories] if plot.categories else []
        series = [
            ChartSeries(
                name=s.name,
                values=[v if v is not None else 0.0 for v in s.values]  # Заменяем None на 0.0
            )
            for s in plot.series
        ]
        return categories, series

    def parse_presentation(self, pptx_path, output_image_dir) -> PresentationInfo:
        prs = Presentation(pptx_path)
        os.makedirs(output_image_dir, exist_ok=True)
        slides_data = []

        for idx_slide, slide in enumerate(prs.slides):
            shapes_data = []

            for idx_shape, shape in enumerate(slide.shapes):
                coords = Coordinates(
                    x=shape.left.pt,
                    y=shape.top.pt,
                    width=shape.width.pt,
                    height=shape.height.pt,
                )

                if shape.has_text_frame and shape.text.strip():
                    first_run = next((run for p in shape.text_frame.paragraphs for run in p.runs), None)
                    font_style = self.parse_font(first_run) if first_run else None

                    shapes_data.append(TextShape(
                        type="text",
                        text=shape.text.strip(),
                        coordinates=coords,
                        font_style=font_style,
                    ))

                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    img_path = os.path.join(output_image_dir, f"slide_{idx_slide + 1}_img_{idx_shape + 1}.png")
                    with open(img_path, "wb") as img_file:
                        img_file.write(shape.image.blob)

                    shapes_data.append(ImageShape(
                        type="image",
                        image_path=img_path,
                        coordinates=coords,
                    ))

                elif shape.has_chart:
                    categories, series = self.parse_chart(shape.chart)
                    shapes_data.append(ChartShape(
                        type="chart",
                        chart_type=str(shape.chart.chart_type),
                        coordinates=coords,
                        categories=categories,
                        series=series,
                    ))

            slides_data.append(Slide(slide_number=idx_slide + 1, shapes=shapes_data))

        return PresentationInfo(
            width=prs.slide_width.pt,
            height=prs.slide_height.pt,
            slides=slides_data,
        )

    def create_presentation_from_template(self, presentation_info: PresentationInfo, output_pptx_path: str):
        prs = Presentation()
        prs.slide_width = Pt(presentation_info.width)
        prs.slide_height = Pt(presentation_info.height)

        for slide_content in presentation_info.slides:
            slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(slide_layout)

            # Установка цвета фона слайда
            bg_color_hex = self.template.get("background_color", "#FFFFFF")
            bg_fill = slide.background.fill
            bg_fill.solid()
            bg_fill.fore_color.rgb = hex_to_rgb(bg_color_hex)

            for shape in slide_content.shapes:
                coords = shape.coordinates

                if isinstance(shape, TextShape):
                    # Добавление текстового блока
                    textbox = slide.shapes.add_textbox(
                        Pt(coords.x), Pt(coords.y), Pt(coords.width), Pt(coords.height)
                    )
                    text_frame = textbox.text_frame
                    text_frame.clear()
                    text_frame.word_wrap = True

                    p = text_frame.add_paragraph()
                    p.text = textwrap.fill(shape.text, width=50)

                    font_style = shape.font_style or FontStyle(**self.template["default_font"])
                    p.font.size = Pt(font_style.size)
                    p.font.name = font_style.name
                    p.font.bold = font_style.bold
                    p.font.italic = font_style.italic
                    p.font.color.rgb = hex_to_rgb(font_style.color)
                    p.alignment = PP_ALIGN.LEFT

                elif isinstance(shape, ImageShape):
                    slide.shapes.add_picture(
                        shape.image_path,
                        Pt(coords.x), Pt(coords.y),
                        width=Pt(coords.width),
                        height=Pt(coords.height)
                    )

                elif isinstance(shape, ChartShape):
                    chart_type_resolved = resolve_chart_type(shape.chart_type)

                    chart_data = CategoryChartData()
                    chart_data.categories = shape.categories
                    for series_data in shape.series:
                        chart_data.add_series(series_data.name, tuple(series_data.values))

                    chart_shape = slide.shapes.add_chart(
                        chart_type_resolved,
                        Pt(coords.x), Pt(coords.y),
                        Pt(coords.width), Pt(coords.height),
                        chart_data
                    )

                    chart = chart_shape.chart
                    for i, series in enumerate(chart.series):
                        series.format.fill.solid()
                        series.format.fill.fore_color.rgb = RGBColor(0, 0, 255)

            prs.save(output_pptx_path)



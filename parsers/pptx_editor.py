from pptx import Presentation
from pptx.dml.color import RGBColor

class PPTXEditor:
    @staticmethod
    def change_background_color(pptx_path: str, output_pptx_path: str, color_hex: str):
        """
        Изменяет цвет заднего фона всех слайдов в презентации.

        :param pptx_path: Путь к исходной презентации.
        :param output_pptx_path: Путь для сохранения изменённой презентации.
        :param color_hex: Цвет в HEX формате (например, '#00FF00').
        """
        prs = Presentation(pptx_path)
        rgb_color = PPTXEditor.hex_to_rgb(color_hex)

        for slide in prs.slides:
            bg_fill = slide.background.fill
            bg_fill.solid()
            bg_fill.fore_color.rgb = rgb_color

        prs.save(output_pptx_path)
        print(f"✅ Задний фон всех слайдов изменён на {color_hex}!")

    @staticmethod
    def change_chart_colors(pptx_path: str, output_pptx_path: str, color_hex: str):
        """
        Изменяет цвет всех графиков в презентации.

        :param pptx_path: Путь к исходной презентации.
        :param output_pptx_path: Путь для сохранения изменённой презентации.
        :param color_hex: Цвет в HEX формате (например, '#FF0000').
        """
        prs = Presentation(pptx_path)
        rgb_color = PPTXEditor.hex_to_rgb(color_hex)

        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_chart:
                    chart = shape.chart
                    for series in chart.series:
                        series.format.fill.solid()
                        series.format.fill.fore_color.rgb = rgb_color

        prs.save(output_pptx_path)
        print(f"✅ Цвет всех графиков изменён на {color_hex}!")

    @staticmethod
    def hex_to_rgb(hex_color: str) -> RGBColor:
        """
        Конвертирует HEX цвет в RGBColor.

        :param hex_color: Цвет в HEX формате (например, '#FF0000').
        :return: Объект RGBColor.
        """
        hex_color = hex_color.lstrip('#')
        return RGBColor(
            int(hex_color[0:2], 16),
            int(hex_color[2:4], 16),
            int(hex_color[4:6], 16)
        )



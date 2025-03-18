from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE

class PPTXEditor:
    @staticmethod
    def change_background_color(pptx_path: str, output_pptx_path: str, color_hex: str):
        """
        Изменяет цвет заднего фона всех слайдов в презентации.

        :param pptx_path: Путь к исходной презентации.
        :param output_pptx_path: Путь для сохранения изменённой презентации.
        :param color_hex: Цвет в HEX формате (например, '#00FF00').
        """
        prs = Presentation(pptx_path)
        rgb_color = PPTXEditor.hex_to_rgb(color_hex)

        for slide in prs.slides:
            bg_fill = slide.background.fill
            bg_fill.solid()
            bg_fill.fore_color.rgb = rgb_color

        prs.save(output_pptx_path)
        print(f"✅ Задний фон всех слайдов изменён на {color_hex}!")

    @staticmethod
    def change_chart_colors(pptx_path: str, output_pptx_path: str, color_hex: str):
        """
        Изменяет цвет всех графиков в презентации.

        :param pptx_path: Путь к исходной презентации.
        :param output_pptx_path: Путь для сохранения изменённой презентации.
        :param color_hex: Цвет в HEX формате (например, '#FF0000').
        """
        prs = Presentation(pptx_path)
        rgb_color = PPTXEditor.hex_to_rgb(color_hex)

        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_chart:
                    chart = shape.chart
                    for series in chart.series:
                        series.format.fill.solid()
                        series.format.fill.fore_color.rgb = rgb_color

        prs.save(output_pptx_path)
        print(f"✅ Цвет всех графиков изменён на {color_hex}!")

    @staticmethod
    def hex_to_rgb(hex_color: str) -> RGBColor:
        """
        Конвертирует HEX цвет в RGBColor.

        :param hex_color: Цвет в HEX формате (например, '#FF0000').
        :return: Объект RGBColor.
        """
        hex_color = hex_color.lstrip('#')
        return RGBColor(
            int(hex_color[0:2], 16),
            int(hex_color[2:4], 16),
            int(hex_color[4:6], 16)
        )
        

def rgb_to_hex(rgb_color: RGBColor):
    return '#{:02x}{:02x}{:02x}'.format(rgb_color[0], rgb_color[1], rgb_color[2])


def hex_to_rgb(hex_color: str):
    hex_color = hex_color.lstrip('#')
    return RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16)
    )

def resolve_chart_type(chart_type_str):
    chart_type_mapping = {
        "COLUMN_CLUSTERED": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "COLUMN_STACKED": XL_CHART_TYPE.COLUMN_STACKED,
        "COLUMN_STACKED_100": XL_CHART_TYPE.COLUMN_STACKED_100,
        "BAR_CLUSTERED": XL_CHART_TYPE.BAR_CLUSTERED,
        "BAR_STACKED": XL_CHART_TYPE.BAR_STACKED,
        "BAR_STACKED_100": XL_CHART_TYPE.BAR_STACKED_100,
        "LINE": XL_CHART_TYPE.LINE,
        "PIE": XL_CHART_TYPE.PIE,
        "SCATTER": XL_CHART_TYPE.XY_SCATTER,
    }
    return chart_type_mapping.get(chart_type_str.split(" ")[0], XL_CHART_TYPE.COLUMN_CLUSTERED)
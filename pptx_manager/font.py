import colorsys

from collections import defaultdict
from collections.abc import Iterator

import numpy as np

from loguru import logger
from PIL import ImageColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.shapes.autoshape import Shape
from pptx.slide import Slides
from pptx.text.text import Font
from pptx.util import Pt

from pptx_manager.models import TextFrameShape, UpdateTextFrameOpts
from pptx_manager.utils import hex_to_rgb


class TextFrameManager:
    def __init__(self):
        self.pres = None

        self.text_frame_shapes = defaultdict(list[TextFrameShape])  # slide_id -> text_frame_shape

    def _get_font_color(self, slide: Slides, font: Font) -> tuple:
        # https://stackoverflow.com/questions/54692768/python-pptx-read-font-color
        font_color = font.color

        try:
            srgb = font_color.rgb
        except AttributeError:
            srgb = None

        if srgb is not None:
            return hex_to_rgb(str(srgb))

        try:
            theme_color = font_color.theme_color
        except AttributeError:
            theme_color = None

        # if the theme color and rgb is None its the black color
        if theme_color is None:
            return 0, 0, 0

        brightness = font_color.brightness

        accent = theme_color.xml_value
        xpath = f'a:themeElements/a:clrScheme/a:{accent}/a:srgbClr/@val'

        slide_master_part = slide.slide_layout.slide_master.part
        theme_part = slide_master_part.part_related_by(RT.THEME)
        theme = parse_xml(theme_part.blob)

        try:
            hex_color = theme.xpath(xpath)[0]
        except IndexError:
            return 0, 0, 0

        srgb = np.array(ImageColor.getcolor(f'#{hex_color}', 'RGB'))

        srgb = srgb / 255
        h, luminance, s = colorsys.rgb_to_hls(*srgb)
        lum_mod = 100000 * (1 - brightness)
        lum_off = 100000 * brightness
        luminance = luminance * (lum_mod / 100000) + (lum_off / 100000)
        srgb = np.array(colorsys.hls_to_rgb(h, luminance, s))
        srgb = (srgb * 255).round(0).astype(int)

        return tuple(srgb)

    def _get_font_size(self, font: Font) -> float:
        font_size = 12

        if font.size is not None:
            font_size = font.size.pt

        return font_size

    def update_text_frame_shape(self, slide_id: int, shape_id: int | None, opts: UpdateTextFrameOpts):
        if opts.text is not None:
            self._update_text_text_frame_shape(slide_id, shape_id, opts.text)

        if opts.color is not None:
            self._update_color_text_frame_shape(slide_id, shape_id, opts.color)

        if opts.italic is not None:
            self._update_italic_text_frame_shape(slide_id, shape_id, opts.italic)

        if opts.underline is not None:
            self._update_underline_text_frame_shape(slide_id, shape_id, opts.underline)

        if opts.bold is not None:
            self._update_bold_text_frame_shape(slide_id, shape_id, opts.bold)

        if opts.size is not None:
            self._update_size_text_frame_shape(slide_id, shape_id, opts.size)

    def _update_text_text_frame_shape(self, slide_id: int, shape_id: int | None, new_text: str) -> None:
        for frame in self._get_frame(slide_id, shape_id):
            frame.text_manager.paragraphs[0].runs[0].text = new_text

    def _update_color_text_frame_shape(
        self, slide_id: int, shape_id: int | None, new_color: tuple[int, int, int]
    ) -> None:
        for frame in self._get_frame(slide_id, shape_id):
            frame.font_manager.color.rgb = RGBColor(new_color[0], new_color[1], new_color[2])

    def _update_bold_text_frame_shape(self, slide_id: int, shape_id: int | None, bold: bool) -> None:
        for frame in self._get_frame(slide_id, shape_id):
            frame.font_manager.bold = bold

    def _update_italic_text_frame_shape(self, slide_id: int, shape_id: int | None, italic: bool) -> None:
        for frame in self._get_frame(slide_id, shape_id):
            frame.font_manager.italic = italic

    def _update_underline_text_frame_shape(self, slide_id: int, shape_id: int | None, underline: bool) -> None:
        for frame in self._get_frame(slide_id, shape_id):
            frame.font_manager.underline = underline

    def _update_size_text_frame_shape(self, slide_id: int, shape_id: int | None, new_size: int) -> None:
        for frame in self._get_frame(slide_id, shape_id):
            frame.font_manager.size = Pt(new_size)

    def _get_frame(self, slide_id: int, shape_id: int | None) -> Iterator[TextFrameShape]:
        frames = self.text_frame_shapes[slide_id]
        for frame in frames:
            if shape_id is not None and shape_id != frame.shape_id:
                continue
            yield frame

    def _create_undefined_font(self, unified_font: Font, base_font: Font, slide: Slides) -> Font:
        unified_font.name = base_font.name
        unified_font.size = Pt(self._get_font_size(base_font))
        unified_font.bold = base_font.bold
        unified_font.italic = base_font.italic
        unified_font.underline = base_font.underline
        unified_font.language_id = base_font.language_id

        unified_font.color.rgb = RGBColor(*self._get_font_color(slide, base_font))

        return unified_font

    def parse_text_frame_shape(
        self, slide_id: int, shape_id: int, slide: Slides, shape: Shape
    ) -> None | TextFrameShape:
        if shape.text == '':
            return None
        try:
            text_frame = shape.text_frame

            first_run_font = None
            full_text_lines = []
            for p_i, paragraph in enumerate(text_frame.paragraphs):
                paragraph_text = ''
                for r_i, run in enumerate(paragraph.runs):
                    if first_run_font is None:
                        first_run_font = run.font
                    paragraph_text += run.text
                full_text_lines.append(paragraph_text)

            merged_text = '\n'.join(full_text_lines)

            if first_run_font is None:
                logger.warning(f'Could not find first font in text frame, text={merged_text}')
                return None

            text_frame.clear()

            paragraph = text_frame.paragraphs[0]
            single_run = paragraph.add_run()
            single_run.text = merged_text
            unified_font = single_run.font

            unified_font = self._create_undefined_font(unified_font, first_run_font, slide)

            font_size = self._get_font_size(unified_font)
            font_color = self._get_font_color(slide, unified_font)
            font_name = unified_font.name
            bold = unified_font.bold
            italic = unified_font.italic
            underline = unified_font.underline

            text = merged_text

        except Exception as e:
            logger.warning(f'error parsing text_frame = {e}, text_shape = {shape.text}')
            return None

        text_frame_shape = TextFrameShape(
            shape_id=shape_id,
            text=text,
            font_name=font_name,
            font_size=font_size,
            bold=bold,
            italic=italic,
            underline=underline,
            color=font_color,
            text_manager=text_frame,
            font_manager=unified_font,
        )

        self.text_frame_shapes[slide_id].append(text_frame_shape)

        return text_frame_shape

    def test(self, source):
        self.pres = Presentation(source)

        slide_id = 1
        for slide in self.pres.slides:
            shape_id = 1
            for shape in slide.shapes:
                if shape.has_text_frame:
                    res = self.parse_text_frame_shape(slide_id, shape_id, slide, shape)
                    if res is not None:
                        shape_id += 1
            slide_id += 1

        logger.debug(f'frames = {self.text_frame_shapes}')

        self.update_text_frame_shape(
            1,
            None,
            UpdateTextFrameOpts(text='тестовая замена текста', color=(0, 0, 0), italic=True, bold=True, size=12),
        )

        self.pres.save('test.pptx')


if __name__ == '__main__':
    fm = TextFrameManager()

    fm.test('../test_sources/test_dit.pptx')

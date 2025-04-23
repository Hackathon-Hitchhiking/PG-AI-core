import os
import subprocess

from io import BytesIO
from tempfile import TemporaryDirectory
from textwrap import dedent

from loguru import logger
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.shapes.autoshape import Shape

from pptx_manager.figure import FigureManager
from pptx_manager.image import ImageManager
from pptx_manager.models import (
    CreateImageFrameOpts,
    CreateShapeOpts,
    CreateTextFrameOpts,
    ShapeType,
    UpdateTextFrameOpts,
)
from pptx_manager.shape import ShapeManager
from pptx_manager.slide import SlideManager
from pptx_manager.table import TableManager
from pptx_manager.text import TextFrameManager


class PPTXManager(
    TextFrameManager,
    ImageManager,
    ShapeManager,
    SlideManager,
    TableManager,
    FigureManager,
):
    def __init__(self, source: str | None, parse_slide_image: bool = False) -> None:
        super().__init__()
        TextFrameManager.__init__(self)
        ImageManager.__init__(self)
        ShapeManager.__init__(self)
        SlideManager.__init__(self)
        TableManager.__init__(self)
        FigureManager.__init__(self)

        self.pres = Presentation(source)

        self.parse_choice = {
            MSO_SHAPE_TYPE.PICTURE: self._parse_image_shape,
            MSO_SHAPE_TYPE.AUTO_SHAPE: self._parse_auto_shape,
            MSO_SHAPE_TYPE.TEXT_BOX: self._parse_text_shape,
            MSO_SHAPE_TYPE.GROUP: self.parse_group_shape,
        }

        self.source = source

        self.slide_count = len(self.pres.slides)

        self.slide_image = {}

        self.parse_presentation()

        if parse_slide_image:
            self.parse_slide_as_images()

    def get_slide_image(self, slide_id: int) -> bytes:
        return self.slide_image[slide_id]

    def parse_slide_as_images(self):
        with TemporaryDirectory() as temp_dir:
            subprocess.run(
                [
                    os.environ.get('LIBREOFFICE_PATH', 'libreoffice'),
                    '--headless',
                    '--convert-to',
                    'pdf',
                    '--outdir',
                    temp_dir,
                    self.source,
                ],
                check=True,
            )

            pptx_filename = os.path.basename(self.source)
            base_name = os.path.splitext(pptx_filename)[0]
            pdf_path = os.path.join(temp_dir, base_name + '.pdf')

            if not os.path.exists(pdf_path):
                raise FileNotFoundError(f'PDF conversion failed; file not found at {pdf_path}')

            pages = convert_from_path(pdf_path, dpi=200)

            for idx, page in enumerate(pages):
                image_buffer = BytesIO()
                page.save(image_buffer, format='PNG')

                self.slide_image[idx + 1] = image_buffer.getvalue()

    def parse_presentation(self) -> None:
        logger.debug(f'Анализ презентации вызван с количеством слайдов = {len(self.pres.slides)}')
        slide_id = 1

        for slide in self.pres.slides:
            self.parse_slide(slide_id, slide)
            for shape in slide.shapes:
                parse_fn = self.parse_choice.get(shape.shape_type)
                if parse_fn is not None:
                    result = parse_fn(slide_id, self.get_shape_count(slide_id), shape)
                    if result is not None:
                        self.increase_shape_count(slide_id, 1)
            slide_id += 1

    def parse_group_shape(self, slide_id: int, shape_id: int, shape: Shape) -> None:
        for group_shape in shape.shapes:
            parse_fn = self.parse_choice.get(group_shape.shape_type)
            if parse_fn is not None:
                result = parse_fn(slide_id, self.get_shape_count(slide_id), group_shape)
                if result is not None:
                    self.increase_shape_count(slide_id, 1)

    def _parse_auto_shape(self, slide_id: int, shape_id: int, shape: Shape) -> None:
        result = self._parse_text_shape(slide_id, shape_id, shape)
        if result is not None:
            self.increase_shape_count(slide_id, 1)
            return

        match shape.auto_shape_type:
            case MSO_AUTO_SHAPE_TYPE.RECTANGLE:
                result = self.parse_figure_shape(slide_id, shape_id, shape)
                if result is not None:
                    self.increase_shape_count(slide_id, 1)
                    return
            case MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
                result = self.parse_figure_shape(slide_id, shape_id, shape)
                if result is not None:
                    self.increase_shape_count(slide_id, 1)
                    return

    def get_json_schema(self) -> dict:
        pres_json = {}
        for slide_id in range(1, self.slide_count + 1):
            text_json = self.get_text_frame_json(slide_id)
            image_json = self.get_image_json(slide_id)
            slide_json = self.get_slide_json(slide_id)
            figure_json = self.get_figure_frame_json(slide_id)

            pres_json[slide_id] = {'text': text_json, 'image': image_json, 'slide': slide_json, 'figure': figure_json}

        return pres_json

    def create_text_shape(self, slide_id: int, opts: CreateTextFrameOpts) -> str:
        """
        Создает новую текстовую фигуру на указанном слайде с заданными параметрами.

        Эта функция добавляет текстовую фигуру в презентацию и применяет указанные параметры форматирования, включая позицию, размер, текстовое содержание и стиль текста.

        Args:
            slide_id (int): ID слайда, на котором будет создана текстовая фигура.
            opts (CreateTextFrameOpts): Объект, содержащий параметры для создания текстовой фигуры.
                -   width (int): Ширина рамки текста.
                -   height (int): Высота рамки текста.
                -   left (int): Позиция рамки текста по оси X.
                -   top (int): Позиция рамки текста по оси Y.
                -   text (str | None, optional): Текстовое содержимое для фигуры.
                -   color (list[int] | None, optional): Цвет текста в формате RGB кортежа.
                -   size (int): Размер шрифта для текста.
                -   bold (bool | None, optional): Установить текст жирным или нет.
                -   italic (bool | None, optional): Установить текст курсивом или нет.
                -   underline (bool | None, optional): Установить подчеркивание текста или нет.

        Returns:
            str: Сообщение о результате операции с подробным описанием созданной фигуры.
        """
        logger.debug(f'Вызов create_text_shape с параметрами slide_id={slide_id}, opts={opts}')
        shape, shape_id = self._add_shape_on_slide(
            slide_id,
            CreateShapeOpts(
                left=opts.left,
                top=opts.top,
                height=opts.height,
                width=opts.width,
                type=ShapeType.TEXT,
            ),
        )

        self._parse_text_shape(slide_id, shape_id, shape)

        self.update_text_frame_shape(
            slide_id,
            shape_id,
            UpdateTextFrameOpts(
                text=dedent(opts.text),
                color=opts.color,
                size=opts.size,
                bold=opts.bold,
                italic=opts.italic,
                underline=opts.underline,
                font_name=opts.font_name,
            ),
        )

        return f'Создана текстовая фигура на слайде {slide_id} (shape_id: {shape_id}) с параметрами: позиция ({opts.left}, {opts.top}), размер ({opts.width}x{opts.height}), текст: "{opts.text}"'

    def delete_text_shape(self, slide_id: int, shape_id: int) -> str:
        """
        Удаляет текстовую фигуру с указанного слайда.

        Эта функция удаляет текстовую фигуру с заданным shape_id со слайда с заданным slide_id. Она пытается удалить фигуру как из внутренней системы отслеживания, так и из самого объекта презентации.

        Args:
            slide_id (int): ID слайда, содержащего текстовую область для удаления.
            shape_id (int): ID фигуры для удаления.

        Returns:
            str: Сообщение о результате операции с информацией об удаленной фигуре.
        """
        logger.debug(f'Вызов delete_text_shape с параметрами slide_id={slide_id}, shape_id={shape_id}')
        shape = self._get_text_frame_shape(slide_id, shape_id)
        self._delete_text_frame_shape(slide_id, shape_id)

        try:
            el = shape.shape_manager.element
            parent = el.getparent()
            parent.remove(el)
        except AttributeError:
            logger.warning(f'Фигура с shape_id={shape_id} не найдена в shape_manager')
            return f'Неизвестный идентификатор фигуры: {shape_id} на слайде {slide_id}'

        return f'Текстовая фигура с ID {shape_id} успешно удалена со слайда {slide_id}'

    def delete_image_shape(self, slide_id: int, shape_id: int) -> str:
        """
        Удаляет фигуру с изображением с указанного слайда.

        Эта функция удаляет фигуру с изображением с заданным shape_id со слайда с заданным slide_id. Она пытается удалить фигуру как из внутренней системы отслеживания, так и из самого объекта презентации.

        Args:
            slide_id (int): ID слайда, содержащего изображение для удаления.
            shape_id (int): ID фигуры с изображением для удаления.

        Returns:
            str: Сообщение о результате операции с информацией об удаленном изображении.
        """
        logger.debug(f'Вызов delete_image_shape с параметрами slide_id={slide_id}, shape_id={shape_id}')
        shape = self._get_image_frame_shape(slide_id, shape_id)
        self._delete_image_frame_shape(slide_id, shape_id)

        try:
            el = shape.shape_manager.element
            parent = el.getparent()
            parent.remove(el)
        except AttributeError:
            logger.warning(f'Фигура с shape_id={shape_id} не найдена в shape_manager')
            return f'Неизвестный идентификатор фигуры с изображением: {shape_id} на слайде {slide_id}'

        return f'Фигура с изображением (ID {shape_id}) успешно удалена со слайда {slide_id}'

    def create_image_shape(self, slide_id: int, opts: CreateImageFrameOpts) -> str:
        """
        Создает новую фигуру с изображением на указанном слайде с заданными параметрами.

        Эта функция добавляет изображение в презентацию и применяет указанные параметры форматирования, включая позицию и размер.

        Args:
            slide_id (int): ID слайда, на котором будет создана фигура с изображением.
            opts (CreateImageFrameOpts): Объект, содержащий параметры для создания фигуры с изображением.
                -   width (int): Ширина рамки изображения.
                -   height (int): Высота рамки изображения.
                -   left (int): Позиция рамки изображения по оси X.
                -   top (int): Позиция рамки изображения по оси Y.
                -   image (bytes): Байты изображения для размещения в фигуре.

        Returns:
            str: Сообщение о результате операции с подробным описанием созданной фигуры с изображением.
        """
        logger.debug(
            f'Вызов create_image_shape с параметрами slide_id={slide_id}, opts={opts.model_dump(exclude={"image"})}'
        )

        shape, shape_id = self._add_image_on_slide(
            slide_id,
            opts.image,
            CreateShapeOpts(
                left=opts.left,
                top=opts.top,
                height=opts.height,
                width=opts.width,
                type=ShapeType.IMAGE,
            ),
        )

        self._parse_image_shape(slide_id, shape_id, shape)

        return f'Создана фигура с изображением на слайде {slide_id} (shape_id: {shape_id}) с параметрами: позиция ({opts.left}, {opts.top}), размер ({opts.width}x{opts.height})'

    def save(self, path: str) -> None:
        self.pres.save(path)


if __name__ == '__main__':
    pr = PPTXManager('../test_data/test_dit.pptx')

    logger.debug(f'figure = {pr.get_all_figure_frame_json()}')

    pr.save('test_create.pptx')

from typing import Optional, List, Dict, Any
from collections import defaultdict
from pptx import Presentation
from pptx.chart.chart import Chart
from pptx.dml.color import ColorFormat, RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_THEME_COLOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.slide import Slide
from pydantic import BaseModel, Field
import colorsys

class ChartSeriesPoint(BaseModel):
    value: Optional[float]
    label: Optional[str]
    color: Optional[tuple[int, int, int]] 
    data_label: Optional[str] 

class Trendline(BaseModel):
    type: str
    equation: Optional[str]
    r_squared: Optional[float] 

class ChartSeries(BaseModel):
    name: str
    values: List[Optional[float]]
    color: Optional[tuple[int, int, int]]
    points: List[ChartSeriesPoint]
    axis_labels: List[str] 
    trendline: Optional[Trendline] 

class ChartAxis(BaseModel):
    title: Optional[str] 
    visible: bool
    font_color: tuple[int, int, int]
    has_major_grid: bool
    has_minor_grid: bool

class ChartLegend(BaseModel):
    position: Optional[str]
    font_color: tuple[int, int, int] 

class ChartData(BaseModel):
    shape_id: int 
    chart_type: str
    title: Optional[str] 
    legend: Optional[ChartLegend] 
    categories: List[str] 
    series: List[ChartSeries] 
    category_axis: ChartAxis 
    value_axis: ChartAxis


class ChartManager:
    def __init__(self):
        self.charts_data = defaultdict(list[ChartData])
        self.chart_objects = defaultdict(dict)  # {slide_id: {shape_id: Chart}}

    def update_series_color(
        self,
        slide_id: int,
        shape_id: int,
        series_name: str,
        new_color: tuple[int, int, int]
    ) -> None:
        """Изменяет цвет указанной серии данных"""
        chart = self.chart_objects.get(slide_id, {}).get(shape_id)
        if not chart:
            raise ValueError("Chart not found")

        for series in chart.series:
            if series.name == series_name:
                fill = series.format.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(*new_color)
                return
        raise ValueError(f"Series '{series_name}' not found")
    

    def recolor_series_by_slide(
        self,
        slide_id: int,
        color: tuple[int, int, int] = (255, 0, 0),
        series_name: Optional[str] = None
    ) -> None:
        """
        Перекрашивает серии на указанном слайде
        :param slide_id: ID целевого слайда
        :param color: RGB цвет (по умолчанию красный)
        :param series_name: Название серии (None = все серии)
        """
        if slide_id not in self.chart_objects:
            raise ValueError(f"Slide {slide_id} not found or has no charts")

        rgb_color = RGBColor(*color)
        
        for shape_id, chart in self.chart_objects[slide_id].items():
            try:
                for series in chart.series:
                    if series_name and series.name != series_name:
                        continue
                        
                    fill = series.format.fill
                    fill.solid()
                    fill.fore_color.rgb = rgb_color
            except Exception as e:
                print(f"Error in chart {shape_id}: {str(e)}")

    def update_chart_title(
        self,
        slide_id: int,
        shape_id: int,
        new_title: str,
        create_if_missing: bool = True
    ) -> None:
        """Обновляет заголовок графика"""
        chart = self.chart_objects.get(slide_id, {}).get(shape_id)
        if not chart:
            raise ValueError("Chart not found")

        if not chart.has_title:
            if create_if_missing:
                chart.has_title = True
            else:
                raise ValueError("Chart has no title")

        chart.chart_title.text_frame.text = new_title

    def update_legend(
        self,
        slide_id: int,
        shape_id: int,
        position: Optional[str] = None,
        font_color: Optional[tuple[int, int, int]] = None
    ) -> None:
        """Обновляет параметры легенды"""
        chart = self.chart_objects.get(slide_id, {}).get(shape_id)
        if not chart:
            raise ValueError("Chart not found")

        if position:
            if not chart.has_legend:
                chart.has_legend = True
            try:
                chart.legend.position = getattr(XL_LEGEND_POSITION, position.upper())
            except AttributeError:
                raise ValueError(f"Invalid legend position: {position}")

        if font_color:
            legend = chart.legend
            legend.font.color.rgb = RGBColor(*font_color)

    def _get_theme_color(self, theme_color: MSO_THEME_COLOR, slide: Slide) -> Optional[str]:
        try:
            master = slide.slide_layout.slide_master
            theme = master.part.related_parts[RT.THEME]._element
            return theme.xpath(f"//a:clrScheme/a:{theme_color.xml_value}/a:srgbClr/@val")[0]
        except (IndexError, KeyError, AttributeError):
            return None

    def _resolve_color(self, color: ColorFormat, slide: Slide) -> Optional[tuple[int, int, int]]:
        if color.type == MSO_COLOR_TYPE.RGB:
            return (color.rgb.r, color.rgb.g, color.rgb.b)
        
        if color.type == MSO_COLOR_TYPE.SCHEME:
            hex_color = self._get_theme_color(color.theme_color, slide)
            if hex_color:
                return self._apply_brightness(hex_color, color.brightness)
        
        return None

    def _apply_brightness(self, hex_color: str, brightness: float) -> tuple[int, int, int]:
        r, g, b = (int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r/255, g/255, b/255)
        new_l = max(0.0, min(1.0, l + brightness))
        r, g, b = colorsys.hls_to_rgb(h, new_l, s)
        return (int(r*255), int(g*255), int(b*255))

    def _parse_axis(self, axis: Any, slide: Slide) -> ChartAxis:
        return ChartAxis(
            title=axis.axis_title.text_frame.text if axis.has_title else None,
            visible=axis.visible,
            font_color=self._resolve_color(axis.tick_labels.font.color, slide) or (0,0,0),
            has_major_grid=axis.has_major_gridlines,
            has_minor_grid=axis.has_minor_gridlines
        )

    def _get_chart_type_name(self, chart: Chart) -> str:
        try:
            return chart.chart_type.name
        except AttributeError:
            return "UNKNOWN"

    def _get_categories(self, chart: Chart) -> List[str]:
        try:
            return [
                str(cat.label)
                for cat in chart.plots[0].categories
                if cat.label is not None
            ]
        except Exception as e:
            print(f"Error getting categories: {e}")
            return []


    def _parse_data_table(self, plot: Any) -> Optional[Dict]:
        """Парсит таблицу данных через XML"""
        try:
            dTable = plot._element.find('.//c:dTable', namespaces=plot._element.nsmap)
            if dTable is None:
                return None
            
            return {
                'show_legend_keys': dTable.get('showKeys', '0') == '1',
            '   show_data_labels': dTable.get('showDataLabels', '0') == '1'
            }
        except Exception as e:
            print(f"Error parsing data table: {e}")
            return None

    def _parse_series(self, chart: Chart, slide: Slide) -> List[ChartSeries]:
        series_list = []
        try:
            for series in chart.series:
            # Парсинг основных данных серии
                series_color = self._resolve_color(series.format.fill.fore_color, slide)
            
                points = []
                for point in series.points:
                # Полный набор данных для точки
                    point_data = ChartSeriesPoint(
                        value=point.data_value,
                        label=point.data_label.text if point.data_label else None,
                        color=self._resolve_color(point.format.fill.fore_color, slide),
                        data_label=point.data_label.text if point.data_label else None
                    )
                    points.append(point_data)
            
            # Парсинг дополнительных свойств
                axis_labels = [
                    str(label.text) 
                    for label in series.data_labels 
                    if label.text
                ] if series.data_labels else []
            
            # Парсинг линий тренда
                trendline = None
                if series.trendlines:
                    trendline = {
                        "type": series.trendlines[0].type.name,
                        "equation": series.trendlines[0].display_equation,
                    "   r_squared": series.trendlines[0].display_r_squared
                    }
            
                series_list.append(ChartSeries(
                    name=series.name,
                    values=series.values,
                    color=series_color,
                    points=points,
                    axis_labels=axis_labels,
                    trendline=trendline
                ))
        except Exception as e:
            print(f"Error parsing series: {e}")
        return series_list

    def parse_chart(self, chart: Chart, slide: Slide, shape_id: int) -> Optional[ChartData]:
        try:
            plot = chart.plots[0]
        
            return ChartData(
                shape_id=shape_id,
                chart_type=self._get_chart_type_name(chart),
                title=chart.chart_title.text_frame.text if chart.has_title else None,
                legend=self._parse_legend(chart.legend, slide) if chart.has_legend else None,
                categories=self._get_categories(chart),
                series=self._parse_series(chart, slide),
                category_axis=self._parse_axis(chart.category_axis, slide),
                value_axis=self._parse_axis(chart.value_axis, slide),
                data_table=self._parse_data_table(plot)
            )
        except Exception as e:
            print(f"Error parsing chart: {e}")
            return None

    def _parse_legend(self, legend: Any, slide: Slide) -> ChartLegend:
        return ChartLegend(
            position=legend.position.name if legend.position else None,
            font_color=self._resolve_color(legend.font.color, slide) or (0,0,0)
        )

def test(file_path: str):
    pres = Presentation(file_path)
    manager = ChartManager()
    
    # Парсинг данных
    for slide in pres.slides:
        chart_count = 1
        for shape in slide.shapes:
            if shape.has_chart:
                print(shape.shape_id)
                chart = shape.chart
                manager.recolor_series_by_slide(
                    slide_id=shape.shape_id,
                    color=(0, 255, 0)
                )
                chart_data = manager.parse_chart(chart, slide, chart_count)
                if chart_data:
                    manager.chart_objects[slide.slide_id][chart_count] = chart
                    chart_count += 1
    
    pres.save("test_sources/modified_presentation.pptx")

if __name__ == "__main__":
    test("test_sources/test.pptx")
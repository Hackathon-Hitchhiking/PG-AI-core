import io

from collections import defaultdict
from collections.abc import Iterator

from PIL import Image as PilImage
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.parts.image import Image
from pptx.shapes.autoshape import Shape

from pptx_manager.models import ImageFrameShape, UpdateImageFrameOpts
from pptx_manager.utils import get_slide_from_shape


def image_to_byte_array(image: Image):
    img_byte_arr = io.BytesIO()
    image.save(img_byte_arr, format=image.format)
    img_byte_arr = img_byte_arr.getvalue()
    return img_byte_arr


class ImageManager:
    def __init__(self):
        self._pres = None

        self.image_frame_shapes = defaultdict(list[ImageFrameShape])  # slide_id -> image_frame_shapes

    def get_image_json(self, slide_id: int):
        return [shape.model_dump(exclude={'blob', 'shape_manager'}) for shape in self.image_frame_shapes[slide_id]]

    def replace_image(self, slide_id: int, shape_id: int | None, new_picture: bytes):
        for shape in self._get_frame(slide_id, shape_id):
            slide = get_slide_from_shape(shape.shape_manager)

            new_shape = slide.shapes.add_picture(
                io.BytesIO(new_picture), shape.left, shape.top, shape.width, shape.height
            )

            old_pic = shape.shape_manager._element
            new_pic = new_shape._element

            old_pic.addnext(new_pic)
            old_pic.getparent().remove(old_pic)

            shape.shape_manager = new_shape
            shape.blob = new_picture

    def update_image_frame_shape(self, slide_id: int, shape_id: int | None, opts: UpdateImageFrameOpts | dict):
        if isinstance(opts, dict):
            opts = UpdateImageFrameOpts(**opts)

        for shape in self._get_frame(slide_id, shape_id):
            if opts.width is not None:
                shape.shape_manager.width = opts.width

            if opts.height is not None:
                shape.shape_manager.height = opts.height

            if opts.left is not None:
                shape.shape_manager.left = opts.left

            if opts.top is not None:
                shape.shape_manager.top = opts.top

    def _get_frame(self, slide_id: int, shape_id: int | None) -> Iterator[ImageFrameShape]:
        frames = self.image_frame_shapes[slide_id]
        for frame in frames:
            if shape_id is not None and shape_id != frame.shape_id:
                continue
            yield frame

    def parse_image_shape(self, slide_id: int, shape_id: int, shape: Shape) -> defaultdict[str, list[ImageFrameShape]]:
        width = shape.width
        height = shape.height
        image = shape.image
        image_bytes = image.blob

        image_frame_shape = ImageFrameShape(
            shape_id=shape_id,
            width=width,
            height=height,
            left=shape.left,
            top=shape.top,
            blob=image_bytes,
            shape_manager=shape,
        )

        self.image_frame_shapes[slide_id].append(image_frame_shape)

        return self.image_frame_shapes

    def test(self, source):
        self._pres = Presentation(source)

        slide_id = 1
        for slide in self._pres.slides:
            shape_id = 1
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    res = self.parse_image_shape(slide_id, shape_id, shape)
                    if res is not None:
                        shape_id += 1
            slide_id += 1

        new_img = PilImage.open('../test_sources/Pr2.jpg')

        self.replace_image(2, None, image_to_byte_array(new_img))

        self.update_image_frame_shape(2, None, {'width': 500_000, 'top': 500})

        self._pres.save('test_new_img.pptx')


if __name__ == '__main__':
    im = ImageManager()

    im.test('../test_sources/test_dit.pptx')

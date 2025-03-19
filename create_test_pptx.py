"""
create_template.py

Generates a minimal 'test_template.pptx' to be used in unit tests.
"""

from pptx import Presentation
from pptx.util import Inches
import os

def create_test_template(path: str) -> None:
    """
    Creates a minimal PowerPoint template with two slides.
    
    Args:
        path (str): The file path to save the generated .pptx template.
    """
    # Create a new blank presentation
    prs = Presentation()

    # --- Slide 1: Title Slide ---
    slide_layout_title = prs.slide_layouts[0]  # Usually the "Title Slide" layout
    slide1 = prs.slides.add_slide(slide_layout_title)

    # Set title
    slide1.shapes.title.text = "Template Title Slide"

    # (Optional) set subtitle if placeholder is available
    # Typically the subtitle is placeholder[1] in a Title Slide layout:
    if len(slide1.placeholders) > 1:
        slide1.placeholders[1].text = "This is a template subtitle"
    textbox = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    textbox.text_frame.text = "Hello from the template!"
    # --- Slide 2: Title and Content ---
    slide_layout_content = prs.slide_layouts[1]  # "Title and Content" layout
    slide2 = prs.slides.add_slide(slide_layout_content)

    # Title placeholder for second slide
    slide2.shapes.title.text = "Template Slide 2"

    # (Optional) Content placeholder is usually placeholders[1]
    if len(slide2.placeholders) > 1:
        body_shape = slide2.placeholders[1]
        body_shape.text = "Bullet 1\nBullet 2\nBullet 3"

    # Save the template
    prs.save(path)


if __name__ == "__main__":
    # Ensure the 'test_sources' directory exists
    os.makedirs("test_sources", exist_ok=True)

    # Create the template in test_sources/test_template.pptx
    template_path = os.path.join("test_sources", "test_template.pptx")
    create_test_template(template_path)

    print(f"Created test template at: {template_path}")